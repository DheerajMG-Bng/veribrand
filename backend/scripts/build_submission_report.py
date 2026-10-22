"""Author the editable submission report using the bundled document runtime.

PDF export is performed separately through a Word-compatible COM interface; render every final PDF page before
delivery. Run from repository root. No environment secrets are read.
"""
import json
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'docs/submission'
QA=ROOT/'.verification/submission-qa'
OUT.mkdir(parents=True,exist_ok=True)
QA.mkdir(parents=True,exist_ok=True)
plans=json.loads((OUT/'round3-designs.json').read_text(encoding='utf-8'))
ai=json.loads((OUT/'evidence/live-ai.json').read_text(encoding='utf-8'))
events=json.loads((OUT/'evidence/live-feature-verification.json').read_text(encoding='utf-8'))

def architecture():
    image=Image.new('RGB',(1700,790),'white')
    draw=ImageDraw.Draw(image)
    font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',27)
    bold=ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf',29)
    boxes=[(30,55,500,200,'Browser UI','React 19 / Vite / Tailwind'),
           (615,55,1085,200,'FastAPI API','Role and workflow guards'),
           (1200,55,1670,200,'SQLite','Versions / findings / audit'),
           (615,315,1085,460,'Review engine','KB metadata rules + cache'),
           (30,575,500,730,'Groq via HTTPX','Budget + three-model fallback'),
           (615,575,1085,730,'Grounding validator','Quotes / codes / replacements'),
           (1200,575,1670,730,'Human decisions','New version / approvals')]
    for x1,y1,x2,y2,title,sub in boxes:
        draw.rounded_rectangle((x1,y1,x2,y2),radius=18,fill='#F2F7F8',outline='#1F2A44',width=3)
        for text,y,f in [(title,y1+37,bold),(sub,y1+89,font)]:
            draw.text(((x1+x2)/2,y),text,fill='#1F2A44',font=f,anchor='mm')
    def arrow(start,end):
        draw.line((start,end),fill='#0E7C86',width=5)
        x,y=end
        if end[0]>start[0]:tri=[(x,y),(x-16,y-9),(x-16,y+9)]
        elif end[1]>start[1]:tri=[(x,y),(x-9,y-16),(x+9,y-16)]
        else:tri=[(x,y),(x+16,y-9),(x+16,y+9)]
        draw.polygon(tri,fill='#0E7C86')
    arrow((500,128),(615,128));arrow((1085,128),(1200,128));arrow((850,200),(850,315))
    arrow((850,460),(850,575));arrow((500,652),(615,652));arrow((1085,652),(1200,652))
    draw.line([(615,390),(265,390),(265,575)],fill='#0E7C86',width=5)
    draw.polygon([(265,575),(256,559),(274,559)],fill='#0E7C86')
    path=OUT/'architecture.png'
    image.save(path)
    return path

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.66)
sec.left_margin=sec.right_margin=Inches(.7)
sec.footer_distance=Inches(.28)
normal=doc.styles['Normal']
normal.font.name='Calibri';normal.font.size=Pt(11);normal.font.color.rgb=RGBColor(0,0,0)
normal.paragraph_format.space_after=Pt(7)
normal.paragraph_format.line_spacing=1.08
for style,size in [('Title',26),('Subtitle',12),('Heading 1',19),('Heading 2',13),('Heading 3',11)]:
    st=doc.styles[style];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_before=Pt(12 if style!='Title' else 0)
    st.paragraph_format.space_after=Pt(7)
    st.paragraph_format.keep_with_next=True
for name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    rpr=doc.styles[name].element.get_or_add_rPr()
    color=rpr.find(qn('w:color'))
    if color is not None:
        for attr in ['themeColor','themeTint','themeShade']:
            color.attrib.pop(qn('w:'+attr),None)
        color.set(qn('w:val'),'000000')
footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=footer.add_run('VeriBrand Round 2  |  ');run.font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.author='Team 27'
doc.core_properties.title='VeriBrand Round 2 Update'
doc.core_properties.subject='Verified hackathon delivery and Round 3 plan'
doc.core_properties.keywords='VeriBrand, Round 2, Team 27'

def p(text,style=None):
    return doc.add_paragraph(text,style)

def h(text,level=1):
    return doc.add_heading(text,level)

def link(label,url):
    paragraph=p(label+': ')
    hyperlink=OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'),paragraph.part.relate_to(url,RT.HYPERLINK,is_external=True))
    run=OxmlElement('w:r');properties=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'0E7C86');properties.append(color)
    run.append(properties);text=OxmlElement('w:t');text.text=url;run.append(text);hyperlink.append(run)
    paragraph._p.append(hyperlink)

def bullet(text):
    paragraph=p(text,'List Bullet')
    paragraph.paragraph_format.space_after=Pt(5)
    return paragraph

def page(title):
    doc.add_page_break();h(title)

def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,width in enumerate(widths):t.columns[i].width=Inches(width)
    all_rows=[headers]+rows
    for ridx,row in enumerate(all_rows):
        cells=t.rows[0].cells if ridx==0 else t.add_row().cells
        for i,text in enumerate(row):
            cell=cells[i];cell.width=Inches(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text=text
            tcpr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            tcpr.append(borders)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'1F2A44' if ridx==0 else 'F3F6F8' if ridx%2 else 'FFFFFF');tcpr.append(shade)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),'85');el.set(qn('w:type'),'dxa');margins.append(el)
            tcpr.append(margins)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after=Pt(2);paragraph.paragraph_format.line_spacing=1.04
                for run in paragraph.runs:
                    run.font.size=Pt(9.5);run.font.bold=ridx==0
                    run.font.color.rgb=RGBColor(255,255,255) if ridx==0 else RGBColor(0,0,0)
        trpr=t.rows[ridx]._tr.get_or_add_trPr();nosplit=OxmlElement('w:cantSplit');trpr.append(nosplit)
        if ridx==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
    p('').paragraph_format.space_after=Pt(0)
    return t

p('VeriBrand Round 2 Update','Title')
p('Team 27  |  DHEERAJ M G and Charan Teja Arangi  |  22 October 2026','Subtitle')
h('Summary',2)
p('VeriBrand is a working text-first brand-governance demo. Creators review content against approved fictional Acme Hydrate knowledge, inspect cited findings and make explicit decisions. New versions pass to a Manager, with Reviewer sign-off for retained medical risks. The live core workflow and an uncached Groq request passed the submission checks.')
link('Live demo','https://veribrand.onrender.com')
link('Repository','https://github.com/DheerajMG-Bng/veribrand')
p('Roles are simulated, not authentication. This shared hackathon deployment is not a production service. Human testing results and Charan\'s contribution remain awaiting input.')
h('Round 1 plan and delivered implementation',2)
p('The plan baseline below is the team-supplied build brief; the original Round 1 document was not independently inspected.')
table(['Plan baseline','Delivered in Round 2','Why it changed'],[
 ['Gemini','Groq via HTTPX with configured fallback','Use the supplied provider; degrade safely'],
 ['Supabase or PostgreSQL','SQLAlchemy 2 and SQLite','Reduce external setup for the demo'],
 ['pgvector retrieval','Complete compact 19-item KB prompt','Small supplied knowledge set'],
 ['Separate React hosting','React build served by FastAPI','One Python deployment and one origin'],
 ['Real user roles','Header-selected demo roles','Finish approval semantics before identity'],
 ['Image and document review','Text review only','Preserve the complete core workflow'],
 ['Rule implementation choice','Metadata rules, grounding and revalidation','Expose sources and reject unsafe wording'],
 ['Demo reliability choice','KB-matched precomputed sample cache','Reduce judging latency and provider use']
],[1.5,2.65,2.65])

page('Completed features')
p('The following were exercised on the public service by HTTP requests or browser use on 22 October 2026. Automated checks are not human usability sessions.')
for text in [
 'Guided demo creates sample 1; sample selection loads and saves a draft.',
 'Rules-only review highlights draft quotes and returns real KB codes; clicking T-02 opens its approved text.',
 'Live AI review returns an AI-suggested unsupported claim, labelled Unverified, needs a human with Accept disabled.',
 'Accept, edit and reject save explicit decisions; reject without a reason is rejected.',
 'Apply decisions creates a new version. Review is required before submitting the latest version, and unapplied fixes block submission.',
 'Version history, manual revisions and plain before/after comparison work.',
 'Comments and return notes persist against content.',
 'Manager approval routes retained medical risks to Reviewer; reasoned sign-off reaches Approved.',
 'Return with comment, revision, re-review and resubmission work end to end.',
 'Audit records the tested actions, decisions, role, time and AI involvement.',
 'Knowledge Base groups 19 items. The live Admin editor saved a title change and restored the original.',
 'Role switching exposes the relevant controls; wrong-role approval returns 403. AI status displays available after loading.',
 'The public API accepts 5,000 characters, rejects 5,001 and returns a friendly review-rate 429 with Retry-After.',
 'FastAPI serves the React app and frontend deep links from one deployment.'
]:bullet(text)
h('Verification scope',2)
p('Browser checks covered the guided demo, sample form, rules review, cited source modal, AI labels, role selector, status chip and admin editor. Public HTTP checks exercised all decisions and approval states, diff, comments, audit, length boundary and rate limiting. Exact evidence and methods are committed under docs/submission/evidence and docs/submission/VERIFICATION.md.')

page('Architecture and technology')
pic=doc.add_picture(str(architecture()),width=Inches(6.8))
pic._inline.docPr.set('descr','React calls same-origin FastAPI. Workflow writes SQLite; review uses KB rules, cache and budgeted Groq, then grounding and human decisions.')
table(['Layer','Implementation'],[
 ['UI','React 19.1.0, React Router 6.30.1, Vite 6.3.5, Tailwind 3.4.17; navy and teal UI'],
 ['API and validation','Python, FastAPI 0.115.12, Pydantic 2.11.4, Uvicorn 0.34.2'],
 ['Data','SQLAlchemy 2.0.40 and SQLite; content, versions, findings, comments, audit, review cache and AI usage'],
 ['Provider','HTTPX 0.28.1 POST to api.groq.com/openai/v1/chat/completions; no provider SDK'],
 ['Model chain','openai/gpt-oss-120b -> openai/gpt-oss-20b -> qwen/qwen3.8-27b -> rules only'],
 ['Deployment','One Render Python service; committed static React assets; health route /api/health']
],[1.25,5.55])
h('API and workflow',2)
p('The /api routes provide knowledge and samples; content creation and reads; review; finding decisions; fix application and versions/diff; submit, approve, return and reviewer-signoff; comments/audit; AI status; and Admin knowledge/reset actions. X-Role selects the simulated role. State-changing operations write audit entries. Approval is an internal state and does not publish externally.')
p('The primary model was verified live during this submission check. The fallback appeared in the original saved run. The configured third model has not been verified live; mocked tests establish fallback control flow, not provider availability.')

page('Technical decisions and difficult parts')
h('Grounding and human control',2)
p('Rules read KB metadata rather than branching on policy codes. The validator maps quotes to actual draft offsets, removes unmatched quotes, nulls invented sources and restricts unsupported findings. Overlap merging preserves the higher-severity issue, supporting codes and Reviewer risk. It does not prove semantic truth merely because a code exists.')
p('Suggested replacement phrases are rechecked for content violations; whole-draft channel disclosure is checked on mandatory re-review. Accepted and edited spans apply from right to left, so offsets stay valid. Nothing is automatically applied, and unapplied decisions cannot be submitted as though they changed the text.')
h('Provider reliability and caching',2)
p('One logical AI analysis can make retries/fallback requests. The client uses a 30-second per-request timeout, JSON output, low temperature and a 2,500 completion-token cap. GPT-OSS uses low reasoning effort, retried once without that field after 400. Short Retry-After is honoured; errors fall through to rules-only results.')
p('Per-model minute usage is in memory and daily usage is in SQLite; local skip thresholds are 6,400 minute tokens and 180,000 daily tokens. SHA-256 cache keys include normalized text, channel, KB hash and AI flag. Precomputed samples preserve genuine model/time provenance. These are single-process controls and do not see other clients using the provider key.')
h('Libraries datasets and licences',2)
table(['Component','Declared licence or provenance'],[
 ['React, React Router, Vite, Tailwind, PostCSS, Autoprefixer','MIT'],
 ['FastAPI, Pydantic, SQLAlchemy, pytest','MIT'],
 ['Uvicorn, HTTPX, python-dotenv','BSD family; see exact package declarations in inventory'],
 ['Lucide icons','ISC'],
 ['Prettier development formatter','MIT'],
 ['Acme Hydrate knowledge and 11 samples','Fictional data supplied in the team build brief; no outside training dataset or fine-tuning'],
 ['Hosted Groq models','Provider/model terms apply separately; configured names do not promise availability']
],[3.05,3.75])
p('Full direct/transitive package names, versions and declared licences are committed in docs/THIRD_PARTY_LICENSES.md. No external website template was copied in this build.')

page('Reliability results and known issues')
table(['Observed check','Result','Scope'],[
 ['pytest rerun','44 passed, 1 warning','Offline rules, validator, workflow and provider mocks'],
 ['Rules sample rerun','17 / 17 expected code occurrences','Our own 11 drafts; clean sample 2 has zero findings'],
 ['AI sample rerun','17 / 17 expected code occurrences','All entries are cached saved responses, revalidated'],
 ['Fresh public AI request',f"0 -> {ai['token_delta']} tokens; cached=false",'Primary GPT-OSS; AI-origin unverified claim'],
 ['Public workflow','Passed','Decisions, versions, approval, return, comments and audit'],
 ['Public boundary checks','Passed','5,000 accepted; 5,001 rejected; friendly 429 observed']
],[1.65,2.05,3.1])
p('Recall figures are measured on our own 11 sample drafts, not general accuracy. They are code-occurrence recall, not independently labelled precision/recall. The original live AI sample run had one false positive on clean sample 2; the general terminology validator correction removed it on replay. Both original and corrected records remain available.')
h('Known limits',2)
for text in [
 'Roles are simulated and data is shared; there is no real authentication or tenant isolation.',
 'Render Free can cold-start and lose SQLite state, audit history and local counters on restart/redeploy/spin-down. Provider quotas still apply.',
 'The third configured model is not verified live. Token estimates are not a preflight guarantee; other uses of the same key are invisible locally.',
 'Confidence is uncalibrated; quote/code validation cannot establish every AI explanation is correct. High-risk overrides can remain with reasons; medical risks need Reviewer sign-off.',
 'Images, DOCX/PDF ingestion, clarity analytics and benchmark metrics in the UI are not implemented.',
 'KB changes do not automatically re-review existing versions or snapshot historical policy text. Diff is plain before/after.',
 'Initial status can briefly say Rules available before loading; rules-only progress copy still mentions AI. The test suite emits a dependency deprecation warning.',
 'No human usability results were supplied. Automated public checks found no blocking core failure; this is not a production or load audit.'
]:bullet(text)

page('Team contributions and disclosure')
h('Team confirmed contributions',2)
p('Dheeraj created the GitHub repository, the Groq key and the Render deployment; wrote the Round 1 plan; directed the AI-assisted build with the prompts; and handled documentation and submission.')
p('The code was generated by Codex (GPT-6.1 Sol) under Dheeraj\'s direction. Codex assisted with tests, verification tooling and submission artifacts. Commit author names are not used to infer who completed a task.')
h('Charan',2)
p('[[HUMAN: Charan to describe only what he actually did]]')
h('User testing',2)
p('[[HUMAN: add results]]')
p('No testers, completion times or participant quotes are claimed. docs/USER_TESTING.md is a ready-to-use sheet with tester initials, role, date, five tasks, time per task, hesitation, bugs and one quote; results are empty. Automated verification is recorded separately.')
h('AI and reuse disclosure',2)
p('Product AI sends draft text, channel, compact approved knowledge and rule summaries to Groq only when enabled on a suitable cache miss. The UI discloses this transfer and asks users to avoid confidential content. The key stays server-side and .env is ignored and untracked. The application uses fictional team-supplied brand data, open-source libraries and Lucide icons; no external website template or additional training dataset was used.')
p('The verified public request returned an AI-origin unsupported claim and increased primary-model daily usage from 0 to 917 tokens. The live UI showed AI-suggested, Unverified, needs a human and disabled Accept. No accounting/label fix was required. An earlier zero alone was inconclusive because cache hits do not add usage and ephemeral counters may reset.')
h('Evidence and source boundaries',2)
p('Implementation and test evidence are in this repository. Round 1 comparisons use the team-supplied plan description; contribution statements use the team\'s explicit confirmation. Round 3 pages describe proposed designs and future acceptance measures, not completed features. Full records: docs/submission/VERIFICATION.md, docs/AI_DISCLOSURE.md and docs/KNOWN_ISSUES.md.')

for i,plan in enumerate(plans,1):
    page('Round 3 improvement '+str(i))
    h(plan['name'].replace('DOCX/PDF','DOCX and PDF').replace('Guideline-Clarity','Guideline Clarity').replace('Multi-format','Multi format'),2)
    p('Proposed work for Round 3. No user impact or acceptance result below is claimed measured.')
    h('User problem',2);p(plan['problem'])
    h('Concrete technical design',2);p(plan['design'])
    h('Milestones',2)
    for m in plan['milestones']:bullet(m)
    h('How success will be measured',2);p(plan['success'])
    h('What we cut first if time is short',2);p(plan['cut'])

# Remove inherited Word title-rule borders as well as any direct paragraph rules.
for root in [doc.element,doc.styles.element]:
    for border in list(root.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.save(OUT/'VeriBrand_Round2_Update.docx')
print('Created editable submission report and architecture figure.')
