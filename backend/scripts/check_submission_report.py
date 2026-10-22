"""Rasterize and structurally check final report using bundled Python libraries.

Visual inspection of every page is still required after this script succeeds.
"""
from pathlib import Path
import json
import re
import pypdfium2 as pdfium
from pypdf import PdfReader
from docx import Document

ROOT=Path(__file__).resolve().parents[2]
QA=ROOT/'.verification/submission-qa'
QA.mkdir(parents=True,exist_ok=True)
pdfpath=ROOT/'docs/submission/VeriBrand_Round2_Update.pdf'
pdf=pdfium.PdfDocument(str(pdfpath))
for i in range(len(pdf)):
    page=pdf[i]
    bitmap=page.render(scale=1.6)
    bitmap.to_pil().save(QA/f'page-{i+1}.png')
    bitmap.close();page.close()
reader=PdfReader(str(pdfpath))
text='\n'.join(page.extract_text() or '' for page in reader.pages)
(QA/'extracted.txt').write_text(text,encoding='utf-8')
doc=Document(ROOT/'docs/submission/VeriBrand_Round2_Update.docx')
paragraphs=[p.text for p in doc.paragraphs if p.text.strip()]
for table in doc.tables:
    paragraphs.extend(p.text for row in table.rows for cell in row.cells for p in cell.paragraphs if p.text.strip())
def norm(value):
    return re.sub(r'\s+','',value).replace('\u00ad','')
missing=[value for value in paragraphs if norm(value) not in norm(text)]
assert not missing,'DOCX text missing from PDF: '+repr(missing[:4])
assert not re.search(r'[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}',text),'Email detected'
assert not re.search(r'(?<!\d)(?:\+91[ -]?)?[6-9]\d{9}(?!\d)',text),'Possible phone number detected'
placeholders=sorted(set(re.findall(r'\[\[HUMAN:.*?\]\]',text)))
assert set(placeholders)=={'[[HUMAN: Charan to describe only what he actually did]]','[[HUMAN: add results]]'}
assert 'https://veribrand.onrender.com' in text
assert 'https://github.com/DheerajMG-Bng/veribrand' in text
links=[annotation.get_object().get('/A',{}).get('/URI') for annotation in reader.pages[0].get('/Annots',[])]
assert 'https://veribrand.onrender.com' in links and 'https://github.com/DheerajMG-Bng/veribrand' in links,'Missing PDF hyperlinks'
report={'pages':len(reader.pages),'paragraphs_and_cells_checked':len(paragraphs),'missing_text':missing,
        'placeholders':placeholders,'first_page_hyperlinks':links,
        'renderer':'Word-compatible COM export (PDF Creator: WPS Docs); pypdfium2 rasterization','visual_review':'pending'}
(QA/'structural-check.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
print(json.dumps(report,indent=2))
